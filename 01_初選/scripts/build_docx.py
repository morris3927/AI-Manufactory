# -*- coding: utf-8 -*-
"""Generate the 初選構想說明書 as a polished .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "PingFang TC"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue
SHADE = "DCE6F1"                       # light blue for layer rows
SHADE2 = "F2F2F2"                      # light grey for arrows/headers

doc = Document()

# ---- default font (incl. CJK) ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
rpr = normal.element.get_or_add_rPr()
rf = OxmlElement("w:rFonts")
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rf.set(qn(a), FONT)
rpr.append(rf)

# ---- margins ----
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.3); s.right_margin = Cm(2.3)


def set_cjk(run, font=FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), color)
    tcpr.append(sh)


def para(text="", size=10.5, bold=False, color=None, align=None,
         before=4, after=4, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        set_cjk(r)
    return p


def h1(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{num}、{text}")
    r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = ACCENT
    set_cjk(r)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), "1F4E79")
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33); set_cjk(r)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.size = Pt(10.5); set_cjk(r)
    r2 = p.add_run(text); r2.font.size = Pt(10.5); set_cjk(r2)
    return p


def table(headers, rows, widths=None, header_shade=SHADE2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        r = c.paragraphs[0].add_run(htext); r.bold = True; r.font.size = Pt(10); set_cjk(r)
        shade(c, header_shade)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].text = ""
            r = cells[i].paragraphs[0].add_run(str(val)); r.font.size = Pt(9.5); set_cjk(r)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def arch_table(steps):
    t = doc.add_table(rows=0, cols=1)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for kind, text in steps:
        c = t.add_row().cells[0]
        c.paragraphs[0].text = ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "arrow" else WD_ALIGN_PARAGRAPH.LEFT
        r = c.paragraphs[0].add_run(text)
        r.font.size = Pt(9.5 if kind == "layer" else 9)
        if kind == "layer":
            r.bold = True; shade(c, SHADE)
        else:
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80); shade(c, "FFFFFF")
        set_cjk(r)
    return t


# ================= COVER / TITLE =================
para("AI Manufacturing 2026 全國競賽 ｜ 主題一 構想說明書", size=11,
     color=RGBColor(0x80, 0x80, 0x80), align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AI 賦能伺服器機殼智慧產線"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = ACCENT; set_cjk(r)
p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("全知辨識與即時響應系統"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor(0x33,0x33,0x33); set_cjk(r)
para("主題一｜AI × 智慧現場：全知辨識與即時響應", size=10.5,
     align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2)
para("團隊名稱：________________　｜　指導老師：________________　｜　成員：________________",
     size=9.5, color=RGBColor(0x80,0x80,0x80), align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=8)

# ================= 一、問題定義 =================
h1("一", "問題定義與應用情境")

h2("1.1 應用場域：伺服器機殼製造廠")
para("本案以伺服器機殼製造為應用情境（為主辦贊助商勤誠興業所屬產業，全球前三大伺服器機殼供應商、NVIDIA 供應鏈）。"
     "其製程涵蓋鈑金沖壓 → 折彎 → 銲接 → 表面處理 → 人工組裝（機殼／電源／線材／背板整合）→ 測試，"
     "屬「金屬加工 + 人工組裝」並重之離散製造。隨 AI 伺服器與液冷需求成長，產線稼動率與工安管理直接牽動交期與良率。")

h2("1.2 現場痛點")
bullet("沖床卡料、機台閒置常延遲數分鐘才被人工巡檢發現，直接損失產能。", "發現慢：")
bullet("異常排查高度依賴資深人員經驗，新人無效嘗試多，平均修復時間（MTTR）偏高。", "處理慢：")
bullet("影像與設備 Log 各看各的，缺乏跨來源關聯——例如機台實際停止、系統卻仍顯示運轉，無人察覺。", "資訊孤島：")
bullet("老師傅經驗屬隱性知識，難以留存與跨廠複製；勤誠正擴建馬來西亞、美國新廠，新廠人員上手慢。", "知識流失：")

h2("1.3 量化現況")
para("離散製造 OEE（設備綜合效率）中位數約 60%（世界級 85%），非計畫停機與偏高 MTTR 是主要缺口；"
     "本系統即針對此兩項缺口設計。")

# ================= 二、解決方案 =================
h1("二", "解決方案構想與 AI 工具技術應用")

h2("2.1 核心主張：決策權分層的自動化閉環")
para("建立「偵測 → 診斷 → 派工 → 現場協作 → 回報 → 知識回寫」之自動化閉環。"
     "設計核心為「決策權分層」：確定性、需秒級反應且要可稽核的事件交給規則引擎（占約 95%）；"
     "只有需要像人一樣推理判斷的環節，才交給 AI／Agent（占約 5%）。"
     "此舉同時兼顧反應速度、可解釋性與成本，是本案最大差異化。")

h2("2.2 產線專屬視覺辨識（不訓練單一通用大模型）")
para("依不同工段部署各自針對性的輕量辨識模型與事件清單，使模型小而準、邊緣端算力低、易維護：")
bullet("沖床運轉／閒置、送料卡料／堵料。", "沖壓段：")
bullet("火花／煙霧、銲接機器人運轉狀態。", "銲接段：")
bullet("人員在崗、PPE（安全帽／護具）、工序合規。", "組裝段：")
para("擴充機制：新增一條產線 = 套用一份「辨識設定檔（Detection Profile）」＋ 增訓該線專屬事件，"
     "不牽動全系統，部署流程標準化。", size=10)

h2("2.3 兩個 AI 介入點（Agent 只放在真正需要判斷處）")
bullet("L2／L3 異常觸發時，LLM 檢索 SOP 知識庫與歷史維修案例（CBR），輸出 top-k 可能根因與對應處置建議，附 SOP 一併打包進派工通知，讓人員不必盲目排查。",
       "① 診斷推論（LLM + RAG／案例式推理）：")
bullet("綁定單一工單與機台，讓檢修人員現場即時問答、引導逐步排查；前述建議無效時動態給出進階／長尾排查指令；結案時將人員回報消化為結構化知識寫回 RAG。",
       "② 任務型維修 Agent：")

h2("2.4 採用之 AI 工具與技術")
table(["環節", "技術／工具"],
      [["視覺辨識", "YOLO 類輕量模型（可邊緣部署）"],
       ["事件判定／告警", "規則引擎 / CEP（複雜事件處理）"],
       ["資料管道", "輕量訊息佇列（正式架構可換 Kafka）"],
       ["知識與診斷", "向量資料庫 RAG + 維修案例庫（CBR）"],
       ["推理與協作", "LLM Agent（診斷推論 + 任務型維修 Agent）"],
       ["通知派工", "Slack / Discord / Email"],
       ["追蹤呈現", "戰情儀表板 + 任務 Checklist"]],
      widths=[3.5, 12])

# ================= 三、系統架構（框架） =================
h1("三", "系統架構、流程圖與應用邏輯")

h2("3.1 分層架構：快路（確定性）× 慢路（Agent）")
arch_table([
    ("layer", "【感測層】影像事件（YOLO，依工段）＋ 設備事件（MES Log 門檻判定）"),
    ("arrow", "↓"),
    ("layer", "【快路 ─ 規則引擎 / CEP（確定性，零 LLM 成本）】嚴重度分級 L1/L2/L3 · 已知跨模態組合告警 · 依職掌路由派工"),
    ("arrow", "↓（僅 L2 / L3 異常啟動 AI；L1 僅記錄）"),
    ("layer", "【AI ① 診斷推論 ─ LLM + RAG / CBR】top-k 可能根因 ＋ 處置建議 ＋ SOP"),
    ("arrow", "↓"),
    ("layer", "【派工 Dispatch】Slack / Discord / Email → 現場人員（附診斷建議）"),
    ("arrow", "↓"),
    ("layer", "【AI ② 任務型維修 Agent】綁定工單＋機台：現場問答 · 引導排查 · 回收檢修邏輯"),
    ("arrow", "↓"),
    ("layer", "【結案閉環】Checklist 確認解除 · 知識回寫 RAG（自我進化）· 戰情儀表板更新"),
])

h2("3.2 嚴重度三級判定")
table(["層級", "定義", "處置"],
      [["L1 Normal", "正常運作", "僅記錄"],
       ["L2 Warning", "需人員處理，無立即停機危險", "啟動 AI 診斷 + 派工"],
       ["L3 Error", "重大異常或經跨模態推論之高風險", "立即緊急通報 + 派工"]],
      widths=[3, 8, 4.5])

h2("3.3 跨模態關聯範例（快路規則，可稽核）")
bullet("影像判斷「機台 Idle／無動作」 vs MES Log「顯示運轉中」→ 兩者不一致 → 感測器或 Log 失準 → 告警。", "核心：")
bullet("影像 Idle ＋ Log Idle 且持續 N 分鐘 → 確認真閒置 → 派工查原因。")
bullet("Log 馬達高溫 ＋ 影像現場無人 → 升級 L3 緊急通報。")

# ================= 四、技術實作深化（實作細節） =================
h1("四", "技術實作深化")

h2("4.1 模型訓練（產線專屬）")
bullet("以預訓練 YOLO 為基礎做遷移學習，各工段分別微調，不從零訓練。", "方法：")
bullet("公開資料集（PPE／工安／表面瑕疵）＋ 現場錄影標註 ＋ 資料增強／合成（模擬卡料、煙霧、低光）。", "資料：")
bullet("蒐集 → 標註 → 微調 → 驗證（mAP／Precision／Recall）→ 部署 → 監控資料漂移 → 迭代。", "流程：")
bullet("每工段一份 Detection Profile（類別清單＋權重檔＋觸發規則）；新事件以小樣本增量訓練，只更新該產線，不動其他線。", "產物與擴充：")

h2("4.2 LLM 選型：線上 vs 地端（資安關鍵決策）")
table(["方案", "優點", "缺點"],
      [["線上雲端 API（如 Claude / GPT 類）", "推理最強、零維運、隨時更新", "資料外送，工廠機密疑慮"],
       ["地端開源模型（如 Llama / Qwen 類，自架 GPU）", "資料不出廠、可控、可離線", "需 GPU 算力、推理略弱"]],
      widths=[6, 5, 4.5])
para("本案建議「混合分層，預設走地端」：含廠內資料的診斷與維修對話走地端開源模型（資料主權）；"
     "僅地端信心不足之長尾複雜案例，經去識別化後選擇性升級雲端。此設計呼應簡章資安要求，"
     "且對勤誠（服務 NVIDIA／Google、NDA 嚴格）最具說服力。", size=10)

h2("4.3 Kafka 資料管道")
bullet("邊緣 CV 節點（影像事件）、MES Log Adapter、門禁系統 Adapter。", "Producers：")
bullet("vision.events / mes.logs / access.events / alerts / tasks。", "Topics：")
bullet("CEP 規則引擎、診斷服務、戰情儀表板、資料封存。", "Consumers：")
bullet("各事件皆帶時間戳，中控以時間窗對齊，作為跨模態關聯依據；原型以輕量 queue／單機 Kafka 等比例縮小，Topic 與 schema 設計不變。", "時間同步與原型：")

h2("4.4 前端畫面與系統樣式")
bullet("即時機台健康度、稼動率、當前告警分佈。", "戰情儀表板：")
bullet("待處理／處理中／已解除三欄；卡片含嚴重度、負責人、計時。", "任務看板（Kanban）：")
bullet("診斷建議（top-k 根因）、SOP、Checklist、維修 Agent 對話視窗。", "任務詳情頁：")
bullet("React／Next.js + WebSocket（即時推播）+ 圖表庫；工業深色「戰情模式」介面。", "技術／樣式：")

h2("4.5 介面串接（整合架構）")
bullet("REST（查詢／CRUD）＋ WebSocket（即時事件推播）。", "後端 API：")
bullet("CV 服務、CEP、診斷（LLM）、維修 Agent、通知、門禁／MES Adapter 全以事件匯流（Kafka）解耦。", "模組解耦：")
bullet("Webhook 串接 Slack／Discord／Email；角色化存取控制（主管／工程師／訪客）。", "通知與權限：")

h2("4.6 機台負責人與門禁整合 ── 責任追溯與精準派工")
bullet("介接產線門禁／刷卡資料，建立「時段 × 區域（機台）× 負責工程師」對應關係。", "資料串接：")
bullet("異常發生時依當下班別與負責人，直接指派該機台當值負責人，而非全體廣播，縮短到場時間。", "精準派工：")
bullet("記錄異常時段之負責人與處理歷程，形成可追溯紀錄，供績效分析與流程改善（正向、非懲罰導向）。", "責任追溯：")
bullet("依簡章資安與個資要求，人員資料採角色化／去識別化處理、存取受權限控管；對外定位為「責任歸屬與可追溯性」。", "個資與治理：")

# ================= 五、預期成果、可行性與效益 =================
h1("五", "預期成果形式、落地可行性與應用效益")

h2("5.1 預期成果形式（決賽 Demo）")
para("可運行系統原型：CV 偵測 → 規則告警 → AI 診斷建議 → 維修 Agent 對話 → 儀表板閉環全流程；"
     "另含系統架構與技術說明、系統測試報告、效益量化分析報告。")

h2("5.2 落地可行性")
para("（1）目標架構 vs 本次原型範圍分層，兼顧遠見與可實現：", size=10)
table(["元件", "目標架構（Production）", "本次原型（Demo）"],
      [["影像來源", "邊緣端輕量過濾、僅推關鍵幀", "公開資料集 / 錄製影片跑 CV"],
       ["訊息佇列", "Kafka 高吞吐", "輕量 queue（Demo 效果相同）"],
       ["診斷推論", "LLM + 完整案例庫", "LLM + 自建 SOP / 案例庫"],
       ["維修 Agent", "多機台多工單併行", "單場景單工單對話"],
       ["資料來源", "真實 MES / 攝影機", "公開資料集 + Log 模擬器"]],
      widths=[2.8, 6.5, 6.2])
bullet("依簡章規定使用公開資料集與模擬資料，不涉及企業機密；影像以公開 PPE／工安資料集，設備 Log 以模擬器產生並與影像事件時間戳對齊。",
       "資料策略：")
bullet("LLM 僅在約 5% 的異常事件觸發（非每幀呼叫），每月僅數十次推論，token 成本極低且可估算。",
       "成本可控：")

h2("5.3 應用效益")
para("產業基準（公開資料，作為合理性背書）：", size=10)
table(["指標", "公開產業基準", "與本系統關聯"],
      [["離散製造 OEE", "中位數約 60%、世界級 85%", "以提升可用率推升 OEE"],
       ["非計畫停機", "即時／預測式維護首年可降 30–50%", "自動偵測 + 引導維修縮短停機"],
       ["MTTR", "導入後降 20–40%", "診斷建議 + 維修 Agent 消除無效排查"]],
      widths=[3, 7.5, 5])
para("假設情境試算（以一條沖壓／銲接產線為例，取值保守、落於上表區間；決賽以 Demo 數據替換）：", size=10, before=6)
para("基準設定：計畫生產時間 16h × 22 天 = 352 h/月；每月停機 60 次，每次 = 發現延遲 12 分（人工巡檢）+ 修復 30 分 = 42 分 → 月停機約 42 h（≈12%）。",
     size=9.5, color=RGBColor(0x66,0x66,0x66))
table(["指標", "現況（人工巡檢）", "導入後", "改善"],
      [["異常平均發現時間", "約 12 分鐘", "約 0.5 分鐘（秒級）", "↓ >95%"],
       ["平均修復時間 MTTR", "30 分鐘", "22.5 分鐘", "↓ 25%"],
       ["每月非計畫停機", "約 42 h", "約 23 h", "↓ 45%"],
       ["設備稼動率 OEE", "約 60%", "約 64%", "↑ 約 4 個百分點"],
       ["每月回收產能", "—", "約 19 h / 線", "直接轉為產出"]],
      widths=[4, 4, 4, 3.5])

h2("5.4 永續與治理價值")
bullet("安全告警走可稽核的規則路徑（非黑箱 LLM），決策可解釋。", "可信賴：")
bullet("維修知識永久留存、隱性知識資產化，支援勤誠馬來西亞／美國新廠快速上線與跨廠經驗傳承。", "知識傳承：")

doc.save("/Users/morris/Desktop/AI_Manufacturing/構想說明書_初選_v1.docx")
print("saved")
