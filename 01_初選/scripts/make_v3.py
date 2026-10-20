# -*- coding: utf-8 -*-
"""v2 -> v3：依組員回饋新增內容，所有新增文字標黃底(yellow highlight)。"""
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/Users/morris/Desktop/AI_Manufacturing/構想說明書_初選_v2.docx"
DST = "/Users/morris/Desktop/AI_Manufacturing/構想說明書_初選_v3.docx"
FONT = "Arial Unicode MS"

d = docx.Document(SRC)
ps = d.paragraphs

# 先抓住要插入位置的元素(插入後 index 會位移，所以先存元素參考)
A_cover   = ps[0]._p     # 封面副標
A_sec2    = ps[21]._p    # 二、 heading -> 後面接系統定義
A_22end   = ps[29]._p    # 2.2 最後一條 -> 後面接視覺辨識小結
P_23note  = ps[31]       # 2.3 診斷推論段(直接在尾巴補註解)
A_img     = ps[36]._p    # 架構圖 -> 後面接 A/B/C/D 說明
A_32head  = ps[37]._p    # 3.2 heading -> 後面接嚴重度前言(在表格之前)
A_33end   = ps[41]._p    # 3.3 最後一條 -> 後面接 3.4 範例
A_54end   = ps[82]._p    # 5.4 最後一條 -> 後面接 SDGs


def style_run(run, bold=False):
    rPr = run._r.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    rPr.append(rf)
    if bold:
        rPr.append(OxmlElement("w:b"))
        rPr.append(OxmlElement("w:bCs"))
    hl = OxmlElement("w:highlight")
    hl.set(qn("w:val"), "yellow")
    rPr.append(hl)


def add_spacing(par, before=80, after=80, line=360):
    pPr = par._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(before)); sp.set(qn("w:after"), str(after))
    sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)


def add_bullet(par, num_id):
    pPr = par._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numPr.append(ilvl)
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id)); numPr.append(nid)
    pPr.insert(0, numPr)


def make_par(runs, style=None, bullet=None):
    """runs: list of (text, bold)。回傳新段落的 w:p 元素(尚未定位)。"""
    par = d.add_paragraph(style=style)            # 先加到文件尾，稍後搬移
    for text, bold in runs:
        style_run(par.add_run(text), bold=bold)
    if style is None:
        add_spacing(par)
    if bullet:
        add_bullet(par, bullet)
    return par._p


def insert_after(anchor_elem, blocks):
    """blocks: list of (runs, style, bullet)。依序插入 anchor 之後。"""
    ref = anchor_elem
    for runs, style, bullet in blocks:
        el = make_par(runs, style=style, bullet=bullet)
        ref.addnext(el)
        ref = el


# ===== (命名) 封面副標 =====
insert_after(A_cover, [
    ([("Sentinel AI ── 智慧產線全知守護與極速響應系統", True),
      ("(通用型智慧產線平台:凡具備監視器影像與機台 Log 之工廠皆可導入)", False)], None, None),
])

# ===== 回饋(1) 二、開頭：系統定義 + 命名 =====
insert_after(A_sec2, [
    ([("為全面破解前述「發現慢、處理慢、資訊孤島、知識流失」四大現場痛點，本團隊打造「", False),
      ("Sentinel AI — 智慧產線全知守護與極速響應系統", True),
      ("」。本系統為通用型智慧產線平台，不限於伺服器機殼產業——只要產線具備「監視器影像」與「機台 Log」兩種訊號即可導入；本案僅以主辦贊助商勤誠興業之伺服器機殼製造作為落地驗證情境。其核心是以「決策權分層」為基礎的自動化閉環:跨模態全知感知現場狀況，由規則引擎提供秒級且可稽核的守護(約 95%)，AI／Agent 僅介入需推理判斷的長尾環節(約 5%)，最終形成可自我進化的知識閉環。", False)], None, None),
])

# ===== 回饋(4-前) 2.2 視覺辨識小結 =====
insert_after(A_22end, [
    ([("小結:", True),
      ("本系統採「一條產線、一個專屬輕量模型」策略，而非訓練單一包山包海的通用大模型。依不同產線各自訓練專屬 AI 影像辨識模型，使每個模型只需辨識該產線的少數事件類別——模型更小、邊緣端運算負擔更低，辨識準確率也更高；新增產線時僅需擴充對應辨識設定檔，不影響既有產線。", False)], None, None),
])

# ===== 回饋(2) 2.3 L2／L3 首次出現處加註解(直接補在原段落尾巴) =====
style_run(P_23note.add_run(
    "(註:此處 L1／L2／L3 為事件嚴重度分級，程度由低到高為 L1 → L2 → L3，完整定義詳見 3.2。)"),
    bold=False)

# ===== 回饋(3) 3.1 規則引擎 A/B/C/D 模組說明(接在架構圖之後) =====
insert_after(A_img, [
    ([("如上方架構圖，L3 規則引擎(CEP)為系統「快路」的決策中樞。它並非單一先後流程，而是由四個", False),
      ("並行", True),
      ("的功能模組組成，每一筆進入的事件會同時經過下列判定:", False)], None, None),
    ([("模組 A｜嚴重度分級:", True),
      ("依事件類型與跨模態判斷結果，將每一筆事件分級為 L1／L2／L3，決定後續處置強度。", False)], None, 1),
    ([("模組 B｜已知跨模態組合告警(可稽核規則):", True),
      ("比對影像事件與 MES Log 等多來源訊號，觸發已知的「組合式」告警(例如 3.3 的影像 Idle 對比 Log 運轉)。此模組專責跨來源關聯，單一事件的嚴重度分級仍由模組 A 完成——因此卡料、火花、PPE 等所有影像事件都不會被遺漏。", False)], None, 1),
    ([("模組 C｜派工路由:", True),
      ("依門禁／班別資料，將事件指派給該機台當值負責人，而非全體廣播，以縮短到場時間。", False)], None, 1),
    ([("模組 D｜分流:", True),
      ("依分級結果決定流向——L1 僅記錄;L2／L3 啟動慢路 AI 診斷並發出派工。", False)], None, 1),
])

# ===== 回饋(2) 3.2 嚴重度三級前言(接在 heading 之後、表格之前) =====
insert_after(A_32head, [
    ([("系統依事件嚴重程度，將所有事件分為三級(程度由低到高:L1 → L2 → L3)，分級結果決定處置方式與是否啟動慢路 AI:", False)], None, None),
])

# ===== 回饋(4) 3.4 事件處理全流程範例 =====
insert_after(A_33end, [
    ([("3.4 事件處理全流程範例(對照架構圖走一遍)", False)], "Heading 2", None),
    ([("為說明上述架構在實際事件中如何運作，以下舉兩類典型事件，沿架構圖走一遍:", False)], None, None),
    ([("範例一｜工安事件(組裝段:人員未戴 PPE)", True)], None, None),
    ([("模組 A:", True), ("影像偵測到人員未戴安全帽 → 判為 L2 Warning(黃燈)。", False)], None, 1),
    ([("模組 C:", True), ("依門禁資料，指派該區域當值負責人。", False)], None, 1),
    ([("模組 D:", True), ("L2 事件 → 觸發派工通知，並啟動 L4 AI 診斷，隨派工一併附上對應 SOP 與處置建議。", False)], None, 1),
    ([("範例二｜生產／工安事件(銲接段:火花、煙霧;沖壓段:卡料、堵料)", True)], None, None),
    ([("模組 A:", True), ("影像偵測到火花／煙霧大作 → 判為 L3 Error(紅燈，重大工安危險)。", False)], None, 1),
    ([("模組 C:", True), ("立即於現場拉響警報，並指派該區所有當值人員到場。", False)], None, 1),
    ([("模組 D:", True), ("L3 事件 → 立即緊急通報並派工(快路秒級反應)，同時將現場狀況送往 L4 的 AI 準備進行診斷。", False)], None, 1),
])

# ===== 回饋(5) 5.4 SDGs 對應 =====
insert_after(A_54end, [
    ([("永續發展目標(SDGs)對應:", True),
      ("本系統亦直接呼應聯合國永續發展目標——", False)], None, None),
    ([("SDG 16(和平、正義與強大機構):", True),
      ("攸關人身安全的工安與即時告警，全數保留在可稽核、透明、非黑箱的規則路徑(而非黑箱 LLM)，決策具高度可解釋性與責任可追溯性;不僅符合資安與個資合規，更能避免 AI 幻覺造成的工安誤判，建立數位時代可信賴的科技治理典範。", False)], None, 3),
    ([("SDG 8(尊嚴勞動與經濟成長)與 SDG 9(產業、創新與基礎建設):", True),
      ("透過維修 Agent 與知識回寫機制，將資深師傅的隱性知識永久留存、轉化為企業結構化資產，有效縮短新進人員的無效排查摸索期(MTTR 降低約 25%)，實現包容性人才培育;此機制更打破地理限制，直接支援勤誠海外(馬來西亞、美國)新廠快速上線與跨廠經驗複製，緩解少子化與跨國擴廠的技術斷層危機。", False)], None, 3),
])

d.save(DST)
print("saved:", DST)
