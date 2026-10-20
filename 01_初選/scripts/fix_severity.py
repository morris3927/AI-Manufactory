# -*- coding: utf-8 -*-
"""把嚴重度 L1/L2/L3 改為 S1/S2/S3(架構層 L0-L6 保留不動)。新/改動文字標黃底。"""
import copy
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/Users/morris/Desktop/AI_Manufacturing/構想說明書_初選_v3.docx"
F = "Arial Unicode MS"
d = docx.Document(PATH)
ps = d.paragraphs


def set_run_text(r_el, text):
    t = r_el.find(qn("w:t"))
    if t is None:
        t = OxmlElement("w:t"); r_el.append(t)
    t.text = text
    t.set(qn("xml:space"), "preserve")


def add_highlight(r_el):
    rPr = r_el.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr"); r_el.insert(0, rPr)
    if rPr.find(qn("w:highlight")) is None:
        hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow"); rPr.append(hl)


def split_highlight(p, old, new):
    """在非黃底段落:把 old 換成 new，並只把 new 標黃底(切 run)。"""
    for run in list(p.runs):
        if old in run.text:
            txt = run.text
            i = txt.index(old)
            pre, post = txt[:i], txt[i + len(old):]
            r = run._r
            r_new = copy.deepcopy(r)
            r_post = copy.deepcopy(r)
            set_run_text(r_new, new); add_highlight(r_new)
            set_run_text(r_post, post)
            run.text = pre
            r.addnext(r_post)
            r.addnext(r_new)
            return True
    return False


def replace_in_yellow(p, pairs):
    """已是黃底的段落:直接在 run 內字串替換(格式維持黃底)。"""
    for run in p.runs:
        for old, new in pairs:
            if old in run.text:
                run.text = run.text.replace(old, new)


# ===== 已黃底段落:直接替換 =====
replace_in_yellow(ps[41], [("L1／L2／L3", "S1／S2／S3")])
replace_in_yellow(ps[44], [("L1", "S1"), ("L2", "S2"), ("L3", "S3")])
replace_in_yellow(ps[46], [("L1", "S1"), ("L2", "S2"), ("L3", "S3")])
replace_in_yellow(ps[54], [("L2 Warning", "S2 警示")])
replace_in_yellow(ps[56], [("L2 事件", "S2 事件")])          # L4 AI 不動
replace_in_yellow(ps[58], [("L3 Error", "S3 異常")])
replace_in_yellow(ps[60], [("L3 事件", "S3 事件")])          # L4 不動

# ===== P34:note(run1,黃底)直接換;原文(run0)切 run 標黃 =====
for run in ps[34].runs:
    if "L1／L2／L3" in run.text:   # 黃底註解
        run.text = run.text.replace("L1／L2／L3", "S1／S2／S3").replace("L1 → L2 → L3", "S1 → S2 → S3")
split_highlight(ps[34], "L2／L3", "S2／S3")   # 原文「L2／L3 異常觸發時」

# ===== P50(3.3,非黃底):升級 L3 → S3 =====
split_highlight(ps[50], "L3", "S3")

# ===== 表格 3.2 第一欄:三個層級標籤 + 標黃底 =====
LABELS = {"L1 Normal": "S1 正常（綠燈）", "L2 Warning": "S2 警示（黃燈）", "L3 Error": "S3 異常（紅燈）"}
for tb in d.tables:
    for row in tb.rows:
        key = row.cells[0].text.strip()
        if key in LABELS:
            cell = row.cells[0]
            p = cell.paragraphs[0]
            for run in p.runs:
                run.text = ""
            run = p.runs[0] if p.runs else p.add_run("")
            set_run_text(run._r, LABELS[key]); add_highlight(run._r)

d.save(PATH)
print("saved")
